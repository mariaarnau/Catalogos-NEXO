from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
wb=Workbook()
F=lambda **k: Font(**{'name':'Arial','size':10,**k})
BLUE=F(color='0000FF'); BOLD=F(bold=True); GREEN=F(color='008000'); NORM=F()
YEL=PatternFill('solid',fgColor='FFFF00'); HEAD=PatternFill('solid',fgColor='DDEBF7')
thin=Side(style='thin',color='999999'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
def hdr(ws,r,vals,c0=1):
    for i,v in enumerate(vals):
        c=ws.cell(r,c0+i,v); c.font=BOLD; c.fill=HEAD; c.border=BOX; c.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center')
def put(ws,ref,v,font=None,fmt=None,note=None):
    c=ws[ref]; c.value=v; c.font=font or (NORM if (isinstance(v,str) and v.startswith('=')) else BLUE if not isinstance(v,str) else NORM)
    if fmt: c.number_format=fmt
    if note: c.comment=Comment(note,'ET48')
    return c
# ---------------- DATOS
d=wb.active; d.title='Datos'
d['A1']='TRABAJO INSTALACIONES DE FLUIDOS 2026-27 · EQUIPO 48 · Hotel 3*'; d['A1'].font=F(bold=True,size=13)
d['A2']='Leyenda: texto azul = dato de entrada (editable) · negro = fórmula · fondo amarillo = hipótesis a confirmar con el tutor'; d['A2'].font=F(italic=True,size=9)
hdr(d,4,['Parámetro','Valor','Unidad','Fuente / nota'])
rows=[
('Equipo',48,'','Enunciado, Anexo I'),
('Tipo de suministro','Mixto (DIR+ALJ)','','Enunciado, Anexo I'),
('Presión de red (a cota calzada 0,00)',40,'mca','Enunciado, Anexo I'),
('Plantas por encima de PB',11,'','Enunciado, Anexo I'),
('Emplazamiento','Barcelona','','Enunciado, Anexo I'),
('Acometidas','Este/Oeste','','Enunciado, Anexo I'),
('Cota calzada',0,'m','Plano esquema vertical'),
('Cota solera PB',0.30,'m','Enunciado cond. 6'),
('Altura PB (solera PB a solera P1)',4.10,'m','Enunciado cond. 6'),
('Altura planta tipo (solera a solera)',3.30,'m','Enunciado cond. 6'),
('Altura conexión aparatos sobre solera',0.60,'m','Enunciado cond. 32'),
('Cota solera sótano',-3.00,'m','HIPÓTESIS: no acotada en el DWG; medir en sección'),
('Mayoración longitudes (f L-eq)',1.30,'','Enunciado cond. 33'),
('Presión mínima en grifo',10,'mca','CTE DB-HS4 2.1.3'),
('Presión máxima en grifo',50,'mca','CTE DB-HS4 2.1.3'),
('Diferencial calderín (BVF) Δp',20,'mca','Procedimiento diseño UPV'),
]
for i,(a,b,c,n) in enumerate(rows):
    r=5+i; d.cell(r,1,a).font=NORM; put(d,f'B{r}',b); d.cell(r,3,c).font=NORM; d.cell(r,4,n).font=F(size=9)
    if 'HIPÓTESIS' in n: d[f'B{r}'].fill=YEL
# UNE 149201
r0=23; d.cell(r0,1,'Simultaneidad UNE 149201 · Hoteles, discotecas, museos:  Qc = A·Qt^B + C  (l/s)').font=BOLD
hdr(d,r0+1,['Rango','A','B','C'])
d.cell(r0+2,1,'Qt ≤ 20 l/s'); put(d,f'B{r0+2}',0.698); put(d,f'C{r0+2}',0.5); put(d,f'D{r0+2}',-0.12)
d.cell(r0+3,1,'Qt > 20 l/s'); put(d,f'B{r0+3}',1.0); put(d,f'C{r0+3}',0.366); put(d,f'D{r0+3}',0)
for rr in (r0+2,r0+3):
    for cc in 'BCD': d[f'{cc}{rr}'].fill=YEL
d[f'B{r0+2}'].comment=Comment('Coeficientes UNE 149201:2008 para hoteles (de memoria del redactor). COMPROBAR con los apuntes / norma antes de entregar.','ET48')
# HS4 table
t0=29; d.cell(t0,1,'Caudales instantáneos mínimos · CTE DB-HS4 Tabla 2.1').font=BOLD
hdr(d,t0+1,['Aparato','Q AF (l/s)','Q ACS (l/s)','Nota'])
ap=[('Lavabo',0.10,0.065,''),('Ducha',0.20,0.10,''),('Bañera ≥ 1,40 m',0.30,0.20,'Bañeras del hotel ≈1,70 m (medido en DWG)'),
('Bidé',0.10,0.065,''),('Inodoro con cisterna',0.10,0,''),('Urinario con cisterna',0.04,0,''),
('Fregadero no doméstico',0.30,0.20,''),('Lavavajillas industrial (20 serv.)',0.25,0.20,'Bitérmico (enunciado cond. 15)'),
('Vertedero',0.20,0,''),('Grifo garaje / limpieza',0.20,0,'')]
for i,(a,b,c,n) in enumerate(ap):
    r=t0+2+i; d.cell(r,1,a).font=NORM; put(d,f'B{r}',b,fmt='0.000'); put(d,f'C{r}',c,fmt='0.000'); d.cell(r,4,n).font=F(size=9)
NAP=len(ap); APR=(t0+2, t0+1+NAP)  # rows 31..40
for col,w in zip('ABCD',(44,16,12,60)): d.column_dimensions[col].width=w
# ---------------- LOCALES
lo=wb.create_sheet('Locales')
lo['A1']='Aparatos por tipo de local (contados sobre el DWG HOTEL_Trabajo_2026-27)'; lo['A1'].font=F(bold=True,size=12)
cols=['Local']+[a[0] for a in ap]+['Qt AF (l/s)','Qt ACS (l/s)','Qt AF+ACS (l/s)','Observaciones']
hdr(lo,3,cols); lo.row_dimensions[3].height=45
locs=[
('Habitación estándar',      [1,0,1,1,1,0,0,0,0,0],'24 por planta tipo, 5 en PB'),
('Habitación adaptada',      [1,1,0,0,1,0,0,0,0,0],'Ducha a ras de suelo con sumidero. 2 por planta tipo, 2 en PB (004, 006)'),
('Oficio de planta (21)',    [0,0,0,0,0,0,1,0,1,0],'Leyenda: fregadero + vertedero'),
('Aseo público hombres (8)', [2,0,0,0,2,2,0,0,0,0],'2 urinarios dibujados en el aseo derecho'),
('Aseo público mujeres (8)', [2,0,0,0,2,0,0,0,0,0],''),
('Vestuario hombres (23)',   [2,1,0,0,1,0,0,0,0,0],'Leyenda: inodoro + ducha + lavabos'),
('Vestuario mujeres (23)',   [2,1,0,0,1,0,0,0,0,0],''),
('Barra (4)',                [0,0,0,0,0,0,1,1,0,0],'Leyenda: fregadero + lavavajillas'),
('Cocina (15)',              [0,0,0,0,0,0,1,2,0,0],'Leyenda: 2 lavavajillas + fregadero'),
('Cuarto frío (17)',         [0,0,0,0,0,0,1,0,0,0],'Leyenda: fregadero'),
('Cuarto de basuras (19)',   [0,0,0,0,0,0,0,0,0,1],'HIPÓTESIS: grifo de limpieza (no dibujado)'),
('Garaje sótano',            [0,0,0,0,0,0,0,0,0,1],'HIPÓTESIS: grifo de limpieza (cond. 21: recogida de aguas en sótano)'),
]
c1=L(2); cN=L(1+NAP)
for i,(n,cnt,obs) in enumerate(locs):
    r=4+i; lo.cell(r,1,n).font=NORM
    for j,v in enumerate(cnt): put(lo,f'{L(2+j)}{r}',v)
    put(lo,f'{L(2+NAP)}{r}','='+'+'.join(f"{L(2+j)}{r}*Datos!$B${APR[0]+j}" for j in range(NAP)),fmt='0.000')
    put(lo,f'{L(3+NAP)}{r}','='+'+'.join(f"{L(2+j)}{r}*Datos!$C${APR[0]+j}" for j in range(NAP)),fmt='0.000')
    put(lo,f'{L(4+NAP)}{r}',f"={L(2+NAP)}{r}+{L(3+NAP)}{r}",fmt='0.000')
    lo.cell(r,5+NAP,obs).font=F(size=9)
    if 'HIPÓTESIS' in obs:
        for j in range(NAP): lo[f'{L(2+j)}{r}'].fill=YEL
NLOC=len(locs); LOR=(4,3+NLOC)
lo.column_dimensions['A'].width=28
for j in range(2,2+NAP+3): lo.column_dimensions[L(j)].width=11
lo.column_dimensions[L(5+NAP)].width=60
open('/tmp/claude-0/-home-user-Catalogos-NEXO/911db6af-93b9-5939-a0d8-cf71bef8bc88/scratchpad/meta.txt','w').write(repr((APR,LOR,NAP,NLOC)))
# ---------------- PLANTAS
pl=wb.create_sheet('Plantas')
pl['A1']='Inventario por planta, caudales y simultaneidad (UNE 149201)'; pl['A1'].font=F(bold=True,size=12)
hcols=['Planta','Nº']+[l[0] for l in locs]+['Cota solera (m)','Cota grifo (m)','Qt AF (l/s)','Qt ACS (l/s)','Qt AF+ACS (l/s)','Qc AF planta (l/s)','Qc ACS planta (l/s)','Qt AF acumulado desde cubierta','Qc AF acumulado (montante) (l/s)']
hdr(pl,3,hcols); pl.row_dimensions[3].height=60
floors=[('Sótano',-1),('Planta baja',0)]+[(f'Planta {k}',k) for k in range(1,12)]
cnt={'Sótano':[0,0,0,0,0,0,0,0,0,0,0,1],'Planta baja':[5,2,1,1,1,1,1,1,1,1,1,0]}
tipo=[24,2,1,0,0,0,0,0,0,0,0,0]
LOC0=3; locL=lambda k: L(LOC0+k)
cz=L(LOC0+NLOC); cg=L(LOC0+NLOC+1); cqa=L(LOC0+NLOC+2); cqs=L(LOC0+NLOC+3); cqt=L(LOC0+NLOC+4); cca=L(LOC0+NLOC+5); ccs=L(LOC0+NLOC+6); cacc=L(LOC0+NLOC+7); cqm=L(LOC0+NLOC+8)
def une(x): return f"=IF({x}<=0,0,MIN({x},IF({x}<=20,Datos!$B$25*{x}^Datos!$C$25+Datos!$D$25,Datos!$B$26*{x}^Datos!$C$26+Datos!$D$26)))"
first=4; last=first+len(floors)-1
for i,(n,k) in enumerate(floors):
    r=first+i; pl.cell(r,1,n).font=NORM; put(pl,f'B{r}',k)
    v=cnt.get(n,tipo)
    for j,x in enumerate(v): put(pl,f'{locL(j)}{r}',x)
    if k==-1: put(pl,f'{cz}{r}','=Datos!$B$16',font=GREEN)
    elif k==0: put(pl,f'{cz}{r}','=Datos!$B$12',font=GREEN)
    else: put(pl,f'{cz}{r}',f'=Datos!$B$12+Datos!$B$13+(B{r}-1)*Datos!$B$14')
    pl[f'{cz}{r}'].number_format='0.00'
    put(pl,f'{cg}{r}',f'={cz}{r}+Datos!$B$15',fmt='0.00')
    rng=f'{locL(0)}{r}:{locL(NLOC-1)}{r}'
    put(pl,f'{cqa}{r}','='+'+'.join(f"{locL(j)}{r}*Locales!${L(2+NAP)}${LOR[0]+j}" for j in range(NLOC)),fmt='0.00')
    put(pl,f'{cqs}{r}','='+'+'.join(f"{locL(j)}{r}*Locales!${L(3+NAP)}${LOR[0]+j}" for j in range(NLOC)),fmt='0.00')
    put(pl,f'{cqt}{r}',f'={cqa}{r}+{cqs}{r}',fmt='0.00')
    put(pl,f'{cca}{r}',une(f'{cqa}{r}'),fmt='0.00')
    put(pl,f'{ccs}{r}',une(f'{cqs}{r}'),fmt='0.00')
    put(pl,f'{cacc}{r}',f'=SUM({cqa}{r}:{cqa}${last})',fmt='0.00')
    put(pl,f'{cqm}{r}',une(f'{cacc}{r}'),fmt='0.00')
tr=last+1
pl.cell(tr,1,'TOTAL EDIFICIO').font=BOLD
for c in (cqa,cqs,cqt): put(pl,f'{c}{tr}',f'=SUM({c}{first}:{c}{last})',font=BOLD,fmt='0.00')
put(pl,f'{cca}{tr}',une(f'{cqa}{tr}'),font=BOLD,fmt='0.00'); put(pl,f'{ccs}{tr}',une(f'{cqs}{tr}'),font=BOLD,fmt='0.00')
pl.cell(tr+1,1,'Qc instalación general común (AF + alimentación producción ACS):').font=BOLD
put(pl,f'{cqt}{tr+1}',une(f'{cqt}{tr}'),font=BOLD,fmt='0.00')
pl.cell(tr+2,1,'Nota: la simultaneidad se aplica al caudal instalado acumulado de cada tramo (UNE 149201). Qc nunca supera Qt.').font=F(italic=True,size=9)
pl.column_dimensions['A'].width=14; pl.column_dimensions['B'].width=5
for j in range(3,3+NLOC+9): pl.column_dimensions[L(j)].width=10
pl.freeze_panes='C4'
# ---------------- ESCALONES
es=wb.create_sheet('Escalones')
es['A1']='Predimensionado: plantas abastecibles en directo con la presión de red'; es['A1'].font=F(bold=True,size=12)
es['A2']='P grifo = P red − (z grifo − z calzada) − Σhm − j·L·fLeq   (Procedimiento diseño UPV, paso 3). Valores de j, Σhm y longitudes = estimación previa al trazado.'; es['A2'].font=F(italic=True,size=9)
inp=[('Pendiente hidráulica de diseño j',40,'mm/m','HIPÓTESIS: rango 15-150 mm/m'),
('Σ pérdidas singulares (llave, filtro, contador, retención, llaves)',5,'mca','HIPÓTESIS: se recalcula con K reales'),
('L acometida → patinillo (incl. bajada a sótano)',20,'m','HIPÓTESIS: acometida por fachada OESTE, núcleo escalera 1 (medir en plano)'),
('L horizontal planta: patinillo → habitación más alejada',47,'m','Medido en DWG: núcleo x≈4 m a hab. extremo x≈51 m'),
('L interior habitación hasta aparato DSF',6,'m','HIPÓTESIS')]
hdr(es,4,['Dato','Valor','Unidad','Nota'])
for i,(a,b,c,n) in enumerate(inp):
    r=5+i; es.cell(r,1,a).font=NORM; put(es,f'B{r}',b); es[f'B{r}'].fill=YEL; es.cell(r,3,c).font=NORM; es.cell(r,4,n).font=F(size=9)
hdr(es,11,['Planta','Cota grifo (m)','L real (m)','hf = j·L·fLeq (mca)','Σhm (mca)','P grifo DSF (mca)','¿Directo? (≥10 mca)'])
for i,(n,k) in enumerate(floors[1:]):
    r=12+i; pr=first+1+i
    put(es,f'A{r}',f'=Plantas!A{pr}',font=GREEN)
    put(es,f'B{r}',f'=Plantas!{cg}{pr}',font=GREEN,fmt='0.00')
    put(es,f'C{r}',f'=$B$7+(B{r}-Datos!$B$16)+$B$8+$B$9',fmt='0.0')
    put(es,f'D{r}',f'=$B$5/1000*C{r}*Datos!$B$17',fmt='0.00')
    put(es,f'E{r}','=$B$6',fmt='0.00')
    put(es,f'F{r}',f'=Datos!$B$7-(B{r}-Datos!$B$11)-D{r}-E{r}',fmt='0.00')
    put(es,f'G{r}',f'=IF(F{r}>=Datos!$B$18,"SÍ","NO → bombeo")')
lr=12+len(floors)-2
es.cell(lr+2,1,'Plantas en directo (incluida PB):').font=BOLD; put(es,f'B{lr+2}',f'=COUNTIF(G12:G{lr},"SÍ")',font=BOLD)
es.cell(lr+3,1,'Plantas con bombeo:').font=BOLD; put(es,f'B{lr+3}',f'=COUNTIF(G12:G{lr},"NO*")',font=BOLD)
es.cell(lr+5,1,'Comprobación nº de plantas por estación con BVF (Procedimiento paso 3): Pmin + Δp calderín + Δz·NP < 50 mca').font=BOLD
es.cell(lr+6,1,'NP máximo con BVF:'); put(es,f'B{lr+6}',f'=INT((Datos!$B$19-Datos!$B$18-Datos!$B$20)/Datos!$B$14)')
es.cell(lr+7,1,'Límite enunciado (cond. 28): BVF ≤ 6 alturas · BVV ≤ 10 alturas por estación').font=F(size=9)
es.column_dimensions['A'].width=58
for c in 'BCDEFG': es.column_dimensions[c].width=15
es.column_dimensions['D'].width=50
wb.save('/home/user/Catalogos-NEXO/Fluidos_ET48/ET48_AF_ACS_calculo_v1.xlsx')
