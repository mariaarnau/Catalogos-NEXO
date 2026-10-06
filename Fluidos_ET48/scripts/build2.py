import sys; sys.path.insert(0,'.')
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
import os; NDIR=int(os.environ.get('NDIR','5')); exec(open('net.py').read()); exec(open('model.py').read().split('print(')[0])
SRC='/home/user/Catalogos-NEXO/Fluidos_ET48/ET48_AF_ACS_calculo_v1.xlsx'
OUT=os.environ.get('OUT','/home/user/Catalogos-NEXO/Fluidos_ET48/ET48_AF_ACS_calculo_v2.xlsx')
wb=load_workbook(SRC)
F=lambda **k: Font(**{'name':'Arial','size':10,**k})
BLUE=F(color='0000FF'); BOLD=F(bold=True); NORM=F(); GREEN=F(color='008000')
YEL=PatternFill('solid',fgColor='FFFF00'); HEAD=PatternFill('solid',fgColor='DDEBF7'); OKF=PatternFill('solid',fgColor='E2EFDA'); BADF=PatternFill('solid',fgColor='F8CBAD')
thin=Side(style='thin',color='999999'); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
def hdr(ws,r,vals,c0=1,h=None):
    for i,v in enumerate(vals):
        c=ws.cell(r,c0+i,v); c.font=BOLD; c.fill=HEAD; c.border=BOX; c.alignment=Alignment(wrap_text=True,horizontal='center',vertical='center')
    if h: ws.row_dimensions[r].height=h
def put(ws,ref,v,font=None,fmt=None,note=None,fill=None):
    c=ws[ref]; c.value=v
    c.font=font or (NORM if (isinstance(v,str) and v.startswith('=')) else (BLUE if not isinstance(v,str) else NORM))
    if fmt: c.number_format=fmt
    if note: c.comment=Comment(note,'ET48')
    if fill: c.fill=fill
    return c
# ---------- Datos: parámetros hidráulicos
d=wb['Datos']
d['A43']='Parámetros hidráulicos y cotas de cálculo'; d['A43'].font=BOLD
hdr(d,44,['Parámetro','Valor','Unidad','Fuente / nota'])
P=[('Gravedad g',9.81,'m/s²',''),
('Viscosidad cinemática AF (20 ºC)',1.004e-6,'m²/s','Propiedades del agua'),
('Viscosidad cinemática ACS (60 ºC)',0.474e-6,'m²/s','Propiedades del agua'),
('Rugosidad absoluta tubos plásticos/multicapa ε',0.007e-3,'m','Valor típico tubería plástica'),
('Cota de la red general (presión de red referida a calzada)',0.0,'m','Enunciado cond. 23'),
('Cota del calderín / colector del grupo de bombeo',-1.2,'m','Cuarto AF en sótano (solera −2,20)'),
('Cota mínima de lámina en aljibe (aspiración)',-2.0,'m','Solera sótano −2,20 + 0,20 resguardo'),
('Cota válvula reductora ACS baja (cuarto ACS)',-1.2,'m','Cuarto ACS en sótano'),
('Pérdidas en el propio grupo (aspiración + valvulería) estimadas',2.0,'mca','HIPÓTESIS: comprobar con modelo de bomba elegido'),
('Rendimiento bomba',0.60,'','Enunciado Anexo II'),
('Rendimiento motor',0.90,'','Enunciado Anexo II'),
('Rendimiento variador',0.95,'','Enunciado Anexo II'),
('Volumen calderín BVV por estación',150,'l','Enunciado cond. 29'),
('Tiempo de reserva aljibe',20,'min','CTE HS4 (15-20 min)'),
]
for i,(a,b,c,n) in enumerate(P):
    r=45+i; d.cell(r,1,a).font=NORM; put(d,f'B{r}',b); d.cell(r,3,c).font=NORM; d.cell(r,4,n).font=F(size=9)
    if 'HIPÓTESIS' in n: d[f'B{r}'].fill=YEL
d['B46'].number_format='0.000E+00'; d['B47'].number_format='0.000E+00'; d['B48'].number_format='0.0E+00'
DG,DNUAF,DNUACS,DEPS,DZRED,DZCAL,DZLAM,DZVRP,DHGR,DETAB,DETAM,DETAV,DVCAL,DTALJ=[f'Datos!$B${45+i}' for i in range(14)]
# K singulares (tabla de referencia)
d['F43']='Coeficientes K de elementos singulares (HIPÓTESIS hasta tener la tabla del profesor)'; d['F43'].font=BOLD
hdr(d,44,['Elemento','K'],c0=6)
KT=[('Llave de corte / aislamiento (bola)',0.3),('Válvula de retención',2.5),('Filtro',5),('Contador',10),('Llave de escuadra en aparato',2),('Acumulador ACS (entrada+salida)',3),('Retención + llave base montante',2.8)]
for i,(a,b) in enumerate(KT):
    d.cell(45+i,6,a).font=NORM; put(d,f'G{45+i}',b,fill=YEL)
d.column_dimensions['F'].width=40; d.column_dimensions['G'].width=8
# ---------- Tubos
tb=wb.create_sheet('Tubos')
tb['A1']='Diámetros comerciales (interior). Multicapa PE-RT/Al/PE-RT PN10 a 70 ºC · PE100 SDR11 PN16 · valvulería por DN'; tb['A1'].font=BOLD
tb['A2']='Comprobar las medidas con el catálogo comercial concreto que elijáis (multicapa: serie habitual 16…90 mm).'; tb['A2'].font=F(italic=True,size=9)
hdr(tb,3,['Serie','Designación','Di (mm)'])
MCr=(4,4+len(MC)-1)
for i,(n,dd) in enumerate(MC): r=4+i; tb.cell(r,1,'MC'); tb.cell(r,2,n); put(tb,f'C{r}',dd)
r0=MCr[1]+3; hdr(tb,r0-1,['Serie','Designación','Di (mm)']); PEr=(r0,r0+len(PE)-1)
for i,(n,dd) in enumerate(PE): r=r0+i; tb.cell(r,1,'PE'); tb.cell(r,2,n); put(tb,f'C{r}',dd)
VAL=[(f'DN{x}',x) for x in (15,20,25,32,40,50,65,80,100,125)]
r0=PEr[1]+3; hdr(tb,r0-1,['Serie','Designación','Di (mm)']); VAr=(r0,r0+len(VAL)-1)
for i,(n,dd) in enumerate(VAL): r=r0+i; tb.cell(r,1,'VAL'); tb.cell(r,2,n); put(tb,f'C{r}',dd)
tb.column_dimensions['B'].width=16
def rng(c,rr): return f'Tubos!${c}${rr[0]}:${c}${rr[1]}'
def nser(rr): return rr[1]-rr[0]+1
# ---------- Tramos
t=wb.create_sheet('Tramos')
t['A1']='Dimensionado de tramos · Qc según UNE 149201 · Darcy-Weisbach (Swamee-Jain) · L cálculo = L real × 1,30'; t['A1'].font=F(bold=True,size=12)
t['A2']='Azul = dato (longitudes medidas en el DWG, nº de locales aguas abajo). Para forzar otro diámetro, cambia la velocidad objetivo (col. L).'; t['A2'].font=F(italic=True,size=9)
H=['ID','Descripción','Circuito','Serie','Hab. estándar aguas abajo','Hab. adaptadas aguas abajo','Oficios aguas abajo','Qt adicional (l/s)','Qt (l/s)','Q aparato mayor (l/s)','Qc (l/s)','v objetivo (m/s)','D teórico (mm)','Designación','Di (mm)','v (m/s)','L real (m)','L cálculo (m)','Re','f','hf (mca)','ΣK','hm (mca)','hf+hm (mca)']
hdr(t,4,H,h=48)
U2=lambda x: f"IF({x}<=0,0,MIN({x},IF({x}<=20,Datos!$B$25*{x}^Datos!$C$25+Datos!$D$25,Datos!$B$26*{x}^Datos!$C$26+Datos!$D$26)))"
for i,(tid,desc,circ,ser,ns,na,no,qx,qmax,vo,L,K) in enumerate(R):
    r=5+i
    put(t,f'A{r}',tid,font=BOLD); t.cell(r,2,desc).font=NORM; t.cell(r,3,circ).font=NORM; t.cell(r,4,ser).font=NORM
    put(t,f'E{r}',ns); put(t,f'F{r}',na); put(t,f'G{r}',no); put(t,f'H{r}',qx,fmt='0.000'); put(t,f'J{r}',qmax,fmt='0.000')
    put(t,f'I{r}',f'=IF(C{r}="AF",E{r}*Locales!$L$4+F{r}*Locales!$L$5+G{r}*Locales!$L$6,E{r}*Locales!$M$4+F{r}*Locales!$M$5+G{r}*Locales!$M$6)+H{r}',fmt='0.000')
    put(t,f'K{r}',f'=MAX(J{r},{U2(f"I{r}")})',fmt='0.000')
    put(t,f'L{r}',vo,fmt='0.0')
    put(t,f'M{r}',f'=SQRT(4*K{r}/1000/(PI()*L{r}))*1000',fmt='0.0')
    def pickf(col):
        return (f'IF(D{r}="MC",INDEX({rng(col,MCr)},MIN(COUNTIF({rng("C",MCr)},"<"&M{r})+1,{nser(MCr)})),'
                f'IF(D{r}="PE",INDEX({rng(col,PEr)},MIN(COUNTIF({rng("C",PEr)},"<"&M{r})+1,{nser(PEr)})),'
                f'INDEX({rng(col,VAr)},MIN(COUNTIF({rng("C",VAr)},"<"&M{r})+1,{nser(VAr)}))))')
    put(t,f'N{r}','='+pickf('B')); put(t,f'O{r}','='+pickf('C'),fmt='0.0')
    put(t,f'P{r}',f'=K{r}/1000/(PI()*(O{r}/1000)^2/4)',fmt='0.00')
    put(t,f'Q{r}',L,fmt='0.00'); put(t,f'R{r}',f'=Q{r}*Datos!$B$17',fmt='0.00')
    put(t,f'S{r}',f'=P{r}*O{r}/1000/IF(C{r}="AF",{DNUAF},{DNUACS})',fmt='0')
    put(t,f'T{r}',f'=IF(S{r}<=0,0,0.25/(LOG10({DEPS}/(3.7*O{r}/1000)+5.74/S{r}^0.9))^2)',fmt='0.0000')
    put(t,f'U{r}',f'=IF(R{r}=0,0,T{r}*R{r}/(O{r}/1000)*P{r}^2/(2*{DG}))',fmt='0.000')
    put(t,f'V{r}',K,fmt='0.0'); put(t,f'W{r}',f'=V{r}*P{r}^2/(2*{DG})',fmt='0.000')
    put(t,f'X{r}',f'=U{r}+W{r}',fmt='0.000')
TR=(5,4+len(R))
widths=[7,52,8,6,9,9,8,10,8,9,8,8,9,12,8,7,8,8,9,8,8,6,8,9]
for j,w in enumerate(widths): t.column_dimensions[chr(65+j)].width=w
t.freeze_panes='C5'
# ---------- Recorridos
rc=wb.create_sheet('Recorridos')
rc['A1']='Recorridos más desfavorables (Bernoulli): P final = P inicio + z inicio − z grifo − Σ(hf+hm)'; rc['A1'].font=F(bold=True,size=12)
rc['A2']='DIR: parte de la red (40 mca a cota 0). BOMB: presión mínima necesaria en el calderín. VRP: presión de tarado necesaria a la salida de la reductora ACS baja.'; rc['A2'].font=F(italic=True,size=9)
hdr(rc,4,['Recorrido','Tipo','Planta','Cota grifo (m)','Σ(hf+hm) (mca)','Resultado (mca)','Interpretación'],h=32)
SUMROW=5; blk=SUMROW+len(PATHS)+3
LOOK=lambda col,ref: f'INDEX(Tramos!${col}${TR[0]}:${col}${TR[1]},MATCH({ref},Tramos!$A${TR[0]}:$A${TR[1]},0))'
for pi,(name,typ,fl,tr) in enumerate(PATHS):
    sr=SUMROW+pi
    hdr(rc,blk,[name,'Descripción','Qc (l/s)','Designación','v (m/s)','hf+hm (mca)','Σ acumulada (mca)'])
    first=blk+1
    for j,tid in enumerate(tr):
        r=first+j
        put(rc,f'A{r}',tid,font=NORM)
        put(rc,f'B{r}','='+LOOK('B',f'A{r}')); put(rc,f'C{r}','='+LOOK('K',f'A{r}'),fmt='0.000')
        put(rc,f'D{r}','='+LOOK('N',f'A{r}')); put(rc,f'E{r}','='+LOOK('P',f'A{r}'),fmt='0.00')
        put(rc,f'F{r}','='+LOOK('X',f'A{r}'),fmt='0.000')
        put(rc,f'G{r}',f'=SUM(F${first}:F{r})',fmt='0.00')
    last=first+len(tr)-1
    put(rc,f'A{sr}',name,font=NORM); rc[f'B{sr}']=typ; put(rc,f'C{sr}',fl)
    put(rc,f'D{sr}',f'=INDEX(Plantas!$P$4:$P$16,MATCH(C{sr},Plantas!$B$4:$B$16,0))',font=GREEN,fmt='0.00')
    put(rc,f'E{sr}',f'=G{last}',fmt='0.00')
    if typ=='DIR':
        put(rc,f'F{sr}',f'=Datos!$B$7+{DZRED}-D{sr}-E{sr}',fmt='0.00')
        put(rc,f'G{sr}',f'=IF(F{sr}>=Datos!$B$18,"Presión en grifo ≥ 10 mca: CUMPLE","NO CUMPLE: pasar a bombeo")')
    elif typ=='BOMB':
        put(rc,f'F{sr}',f'=Datos!$B$18+D{sr}-{DZCAL}+E{sr}',fmt='0.00')
        put(rc,f'G{sr}','Presión mínima necesaria en calderín')
    else:
        put(rc,f'F{sr}',f'=Datos!$B$18+D{sr}-{DZVRP}+E{sr}',fmt='0.00')
        put(rc,f'G{sr}','Presión mínima de tarado de la reductora ACS baja')
    blk=last+3
for c,w in zip('ABCDEFG',(40,46,10,14,9,13,44)): rc.column_dimensions[c].width=w
NP=len(PATHS); PR=(SUMROW,SUMROW+NP-1)
# ---------- Bombeo
bo=wb.create_sheet('Bombeo')
bo['A1']='Grupo de presión (BVV) aspirando de aljibe · Válvula reductora ACS baja · Aljibe'; bo['A1'].font=F(bold=True,size=12)
hdr(bo,3,['Concepto','Valor','Unidad','Cálculo / nota'])
def sumif_max(t):
    rows=[f'Recorridos!$F${PR[0]+i}' for i,p in enumerate(PATHS) if p[1]==t]
    return '=MAX('+','.join(rows)+')'
rows=[
('Presión mínima necesaria en calderín (máx. recorridos BOMB)',sumif_max('BOMB'),'mca','Máximo de los recorridos R5-R12'),
('Presión a mantener (consigna BVV), redondeada',"=CEILING(B4,1)",'mca','Redondeo al mca superior'),
('Caudal de bombeo Qb (tramo BV0)',"="+LOOK('K','"BV0"'),'l/s','AF plantas 6-11 + alimentación ACS de todo el edificio, simultaneidad UNE'),
('Caudal de bombeo Qb',"=B6*3.6",'m³/h',''),
('Nº de bombas en servicio',"=IF(B6<3,1,IF(B6<10,2,IF(B6<30,3,4)))",'','CTE / apuntes: <3 l/s:1 · 3-10:2 · 10-30:3 · >30:4 (+1 reserva)'),
('Nº total de bombas (incl. reserva)',"=B8+1",'',''),
('Caudal por bomba',"=B6/B8",'l/s',''),
('Altura de bombeo H',f"=B5+({DZCAL}-{DZLAM})+{DHGR}",'mca','H = P consigna + (z calderín − z lámina mín.) + pérdidas en grupo'),
('Potencia hidráulica por bomba',f"=1000*{DG}*B10/1000*B11/1000",'kW','ρ·g·Q·H'),
('Potencia en el eje por bomba',f"=B12/{DETAB}",'kW','η bomba'),
('Potencia eléctrica por bomba (motor + variador)',f"=B13/({DETAM}*{DETAV})",'kW','η motor · η variador'),
('Volumen calderín (BVV)',f"={DVCAL}",'l','Enunciado cond. 29'),
('Presión estática máx. AF/ACS bombeo, grifo planta '+str(NDIR+1),f"=B5+{DZCAL}-INDEX(Plantas!$P$4:$P$16,MATCH({NDIR+1},Plantas!$B$4:$B$16,0))",'mca','Q=0: consigna + z cald − z grifo planta más baja con bombeo. Debe ser < 50'),
('Comprobación presión máxima AF/ACS alta',"=IF(B16<Datos!$B$19,\"CUMPLE (<50 mca)\",\"NO CUMPLE: reductoras\")",'',''),
('Tarado mínimo reductora ACS baja (máx. recorridos VRP)',sumif_max('VRP'),'mca','Máximo de R13-R16, R18'),
('Tarado reductora ACS baja adoptado',"=CEILING(B18,1)",'mca',''),
('Presión estática máx. ACS baja, grifo PB',f"=B19+{DZVRP}-INDEX(Plantas!$P$4:$P$16,MATCH(0,Plantas!$B$4:$B$16,0))",'mca','Debe ser < 50'),
('Comprobación presión máxima ACS baja',"=IF(B20<Datos!$B$19,\"CUMPLE (<50 mca)\",\"NO CUMPLE\")",'',''),
('Presión estática máx. AF directo, grifo PB',f"=Datos!$B$7+{DZRED}-INDEX(Plantas!$P$4:$P$16,MATCH(0,Plantas!$B$4:$B$16,0))",'mca','Red sin consumo'),
('Comprobación presión máxima AF directo',"=IF(B22<Datos!$B$19,\"CUMPLE (<50 mca)\",\"NO CUMPLE\")",'',''),
('ALJIBE · volumen necesario',f"=B6*60*{DTALJ}",'l','V = Qb (l/min) × t'),
('ALJIBE · nº de vasos',2,'','Mínimo 2 vasos en paralelo (enunciado)'),
('ALJIBE · volumen necesario por vaso',"=B24/B25",'l',''),
('ALJIBE · vaso comercial adoptado',4000,'l','Catálogo (Cuestiones AF): 4000 l, Ø2,12 × H1,465 m, 4090 kg lleno'),
('ALJIBE · volumen total instalado',"=B25*B27",'l',''),
('ALJIBE · comprobación',"=IF(B28>=B24,\"CUMPLE\",\"AUMENTAR VASO\")",'',''),
]
for i,(a,b,c,n) in enumerate(rows):
    r=4+i; bo.cell(r,1,a).font=NORM; put(bo,f'B{r}',b,fmt='0.00'); bo.cell(r,3,c).font=NORM; bo.cell(r,4,n).font=F(size=9)
for c,w in zip('ABCD',(56,14,8,70)): bo.column_dimensions[c].width=w
wb.calculation.fullCalcOnLoad=True
wb.save(OUT); print('saved',OUT, 'tramos',TR,'paths',PR)
# ================= reabrir y añadir ACS, Resumen y nota en Escalones
wb=load_workbook(OUT)
es=wb['Escalones']
es['A3']=f'NOTA: este es el PREDIMENSIONADO inicial (paso 3 del procedimiento, con j y Σhm estimadas). El cálculo definitivo con el trazado real está en la hoja Recorridos: directo PB–P{NDIR}, bombeo P{NDIR+1}–P11.'
es['A3'].font=F(bold=True,color='C00000')
# ---------- ACS (producción)
ac=wb.create_sheet('ACS')
ac['A1']='Producción centralizada de ACS con aerotermia · Justificación CTE DB-HE4'; ac['A1'].font=F(bold=True,size=12)
hdr(ac,3,['Concepto','Valor','Unidad','Fuente / nota'])
A=[
('Nº de habitaciones del hotel','=SUMPRODUCT(Plantas!C4:C16+Plantas!D4:D16)','ud','Plantas (estándar + adaptadas)'),
('Ocupación por habitación',2,'pers/hab','HIPÓTESIS: habitaciones dobles (2 camas o cama de matrimonio en el DWG)'),
('Nº de personas','=B4*B5','pers',''),
('Demanda unitaria ACS a 60 ºC, hotel ***',41,'l/(pers·día)','CTE DB-HE4 Anejo F (hotel 3*) — COMPROBAR en el CTE vigente'),
('Demanda diaria de ACS a 60 ºC, D','=B6*B7','l/día',''),
('Contribución renovable mínima exigida','=IF(B8>=5000,0.7,0.6)','','HE4: 70 % si D ≥ 5.000 l/día; 60 % si < 5.000 (apuntes ACS)'),
('SCOP mínimo de la bomba de calor para alcanzar esa contribución','=MAX(2.5,1/(1-B9))','','Erenov/Etot = 1 − 1/SCOP (apuntes ACS, BdC); SCOP ≥ 2,5 para ser renovable'),
('SCOP de la bomba de calor elegida',3.6,'','HIPÓTESIS: tomar el valor del catálogo del equipo (BdC alta temperatura)'),
('Contribución renovable obtenida','=1-1/B11','',''),
('Comprobación HE4','=IF(B12>=B9,"CUMPLE","NO CUMPLE: subir SCOP o añadir solar térmica")','',''),
('Temperatura agua fría de red (diseño, invierno)',10,'ºC','HIPÓTESIS: Barcelona, mes más frío (UNE 94002) — comprobar'),
('Energía diaria necesaria','=B8*4.186*(60-B14)/3600','kWh/día','E = V·cp·(60 − T AF)'),
('Horas de funcionamiento de la BdC',16,'h/día','Criterio de diseño adoptado'),
('Potencia térmica bombas de calor','=B15/B16','kW','P = E / t'),
('Volumen de acumulación a 60 ºC (criterio: 40 % de D)','=0.4*B8','l','Criterio de diseño adoptado (cubrir punta mañana/noche)'),
('Nº de acumuladores en paralelo',4,'ud','Depósitos en sótano (enunciado cond. 41)'),
('Volumen por acumulador','=B18/B19','l','Elegir el comercial superior'),
('Caldera de gas natural de apoyo (100 % respaldo)','=B17','kW','Enunciado cond. 39: apoyo si la BdC no alcanza 60 ºC'),
('Retorno de ACS','Sí','','Distancia > 15 m desde la producción (CTE HS4); solo se dibuja (Anexo II)'),
]
for i,(a,b,c,n) in enumerate(A):
    r=4+i; ac.cell(r,1,a).font=NORM; put(ac,f'B{r}',b,fmt='0.00'); ac.cell(r,3,c).font=NORM; ac.cell(r,4,n).font=F(size=9)
    if 'HIPÓTESIS' in n or 'COMPROBAR' in n: ac[f'B{r}'].fill=YEL
ac['B9'].number_format='0%'; ac['B12'].number_format='0.0%'
for c,w in zip('ABCD',(58,14,12,80)): ac.column_dimensions[c].width=w
# ---------- Resumen
rs=wb.create_sheet('Resumen',0)
rs['A1']='RESUMEN · EQUIPO 48 · Entrega 1 (AF + ACS) · Hotel 3* Barcelona'; rs['A1'].font=F(bold=True,size=13)
rs['A2']='Todos los valores son fórmulas enlazadas a las hojas de cálculo.'; rs['A2'].font=F(italic=True,size=9)
hdr(rs,4,['Concepto','Valor','Unidad / detalle'])
LT=lambda col,i: '='+LOOK(col,f'"{i}"')
S=[
('Suministro','Mixto DIR+ALJ',''),
('Plantas en directo',f'PB a P{NDIR}',f'Recorridos R1-R4: presión mínima en grifo P{NDIR}'),
('Presión mínima en grifo más desfavorable en directo','=MIN(Recorridos!F5:F8)','mca (≥10)'),
('Plantas con bombeo',f'P{NDIR+1} a P11','1 estación de bombeo BVV aspirando de aljibe'),
('Caudal simultáneo instalación general (Qc)',LT('K','G1'),'l/s'),
('Acometida / tubo de alimentación',LT('N','G1'),'PE100 SDR11 PN16'),
('Contador general',LT('N','E1'),'Woltman (velocidad 0,8-1,2 m/s)'),
('Válvula de retención general',LT('N','E2'),''),
('Filtro general',LT('N','E4'),''),
('Montante directo O (tramo inicial)',LT('N','MDL1'),'multicapa'),
('Montante bombeo O (tramo inicial)',LT('N',f'MBL{NDIR+1}'),'multicapa'),
('Montante ACS alta O (tramo inicial)',LT('N',f'MCAL{NDIR+1}'),'multicapa'),
('Montante ACS baja O (tramo inicial)',LT('N','MCBL1'),'multicapa'),
('Presión a mantener en calderín (BVV)','=Bombeo!B5','mca'),
('Caudal de bombeo','=Bombeo!B6','l/s'),
('Nº de bombas','=Bombeo!B8&" + 1 reserva"',''),
('Altura de bombeo','=Bombeo!B11','mca'),
('Potencia eléctrica por bomba','=Bombeo!B14','kW'),
('Calderín','=Bombeo!B15','l'),
('Presión estática máxima (bombeo, planta más baja)','=Bombeo!B16','mca (<50)'),
('Tarado válvula reductora ACS baja','=Bombeo!B19','mca'),
('Aljibe','=Bombeo!B25&" vasos × "&Bombeo!B27&" l"','=Bombeo!B29'),
('Demanda ACS','=ACS!B8','l/día a 60 ºC'),
('Contribución renovable exigida / obtenida','=TEXT(ACS!B9,"0%")&" / "&TEXT(ACS!B12,"0%")','=ACS!B13'),
('Potencia bombas de calor','=ACS!B17','kW'),
('Volumen de acumulación ACS','=ACS!B18','l'),
]
for i,(a,b,c) in enumerate(S):
    r=5+i; rs.cell(r,1,a).font=NORM; put(rs,f'B{r}',b,fmt='0.00'); put(rs,f'C{r}',c,font=NORM)
rs['A33']='Hipótesis pendientes de confirmar (celdas amarillas): coeficientes UNE hotel (Datos), K de elementos singulares (Datos), SCOP y T agua fría (ACS), ocupación 2 pers/hab.'
rs['A33'].font=F(bold=True,color='C00000')
for c,w in zip('ABC',(52,22,50)): rs.column_dimensions[c].width=w
wb.calculation.fullCalcOnLoad=True
wb.save(OUT); print('final saved')
